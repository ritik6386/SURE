"""
Document Parser & Procurement Tender Analyzer Service
Extracts technical requirements, Bill of Quantities (BoQ) items, and equipment
specifications from uploaded PDF, DOCX, and TXT tender/DPR documents.
"""

import io
import re
import pypdf
import docx
from typing import List, Dict, Any, Optional

from backend.app.services.recommendation_service import RecommendationEngine

class DocumentParserService:
    def __init__(self, engine: Optional[RecommendationEngine] = None):
        self.engine = engine if engine else RecommendationEngine()

        # Equipment & technical requirement anchor keywords for detection
        self.equipment_keywords = [
            "motor", "pump", "transformer", "cable", "conductor", "wire",
            "switchgear", "circuit breaker", "mcb", "mccb", "isolator",
            "meter", "energy meter", "lighting", "luminaire", "led",
            "solar", "pv module", "inverter", "battery", "ups", "cctv",
            "camera", "computer", "laptop", "server", "fire extinguisher",
            "cement", "steel", "rebar", "pipe", "valve", "fitting",
            "geotextile", "bitumen", "chlorine", "caustic soda", "alum",
            "syringe", "glove", "mask", "catheter", "water", "fertilizer",
            "diesel generator", "dg set", "panel", "distribution board"
        ]

    def extract_text(self, file_bytes: bytes, filename: str) -> str:
        """Extract plain text from uploaded file based on its extension."""
        ext = filename.lower().split(".")[-1]
        text = ""

        if ext == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_content = page.extract_text() or ""
                pages_text.append(f"--- PAGE {i+1} ---\n" + page_content)
            text = "\n".join(pages_text)

        elif ext in ["docx", "doc"]:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs)

        elif ext in ["txt", "csv", "md"]:
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Supported formats: .pdf, .docx, .txt")

        return text.strip()

    def segment_technical_sections(self, raw_text: str) -> List[str]:
        """
        Public procurement tenders contain legal boilerplate (EMD, penalties, arbitration).
        This method identifies technical schedules, BoQ, and specification sections.
        """
        lines = raw_text.splitlines()
        candidate_chunks = []
        current_chunk = []
        in_technical_section = False

        tech_header_pattern = re.compile(
            r"(technical\s+specification|bill\s+of\s+quantit|boq|schedule\s+of\s+requirement|scope\s+of\s+work|equipment\s+schedule|item\s+description|material\s+specification)",
            re.IGNORECASE
        )

        non_tech_header_pattern = re.compile(
            r"(general\s+conditions\s+of\s+contract|gcc|arbitration|force\s+majeure|earnest\s+money|emd|tender\s+fee|commercial\s+terms)",
            re.IGNORECASE
        )

        doc_header_pattern = re.compile(
            r"^(?:government\s+of|notice\s+inviting\s+tender|nit\s+no|tender\s+no|section\s+[ivx0-9]+|part\s+[ivx0-9]+|chapter\s+[ivx0-9]+|project\s*:|subject\s*:)",
            re.IGNORECASE
        )

        has_any_tech_header = any(tech_header_pattern.search(l) for l in lines)

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            if tech_header_pattern.search(line_str):
                in_technical_section = True
                continue
            elif non_tech_header_pattern.search(line_str):
                in_technical_section = False

            if doc_header_pattern.search(line_str):
                continue

            has_equip = any(re.search(rf"\b{kw}\b", line_str, re.IGNORECASE) for kw in self.equipment_keywords)
            has_is_code = bool(re.search(r"\bIS\s*:?\s*\d+", line_str, re.IGNORECASE))

            if in_technical_section or (not has_any_tech_header and (has_equip or has_is_code)):
                current_chunk.append(line_str)
            elif current_chunk:
                if len(current_chunk) >= 1:
                    candidate_chunks.append("\n".join(current_chunk))
                current_chunk = []

        if current_chunk:
            candidate_chunks.append("\n".join(current_chunk))

        if not candidate_chunks:
            candidate_chunks = [raw_text]

        return candidate_chunks

    def extract_discrete_items(self, text: str) -> List[Dict[str, Any]]:
        """
        Decomposes document text into discrete procurement line items.
        Handles numbered lists, bullet points, BoQ rows, or multi-line clauses.
        """
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        items = []
        seen_texts = set()

        doc_header_pattern = re.compile(
            r"^(?:government\s+of|notice\s+inviting\s+tender|nit\s+no|tender\s+no|section\s+[ivx0-9]+|part\s+[ivx0-9]+|chapter\s+[ivx0-9]+|project\s*:|subject\s*:)",
            re.IGNORECASE
        )

        item_prefix_pattern = re.compile(
            r"^(?:item\s*(?:no\.?|#)?\s*\d+[:\.\-]?|sl\.?\s*(?:no\.?|#)?\s*\d+[:\.\-]?|\d+\s*[\.\):]|[-•*])\s*",
            re.IGNORECASE
        )

        current_item = None

        for line in lines:
            if line.startswith("--- PAGE") or doc_header_pattern.search(line):
                continue

            is_new_item_prefix = bool(item_prefix_pattern.match(line))
            has_equip = any(re.search(rf"\b{kw}\b", line, re.IGNORECASE) for kw in self.equipment_keywords)
            has_is = bool(re.search(r"\bIS\s*:?\s*\d+", line, re.IGNORECASE))

            if (is_new_item_prefix and (has_equip or has_is or len(line) > 20)) or (has_equip and not current_item):
                if current_item and len(current_item["text"]) > 10:
                    clean_txt = current_item["text"].strip()
                    if clean_txt not in seen_texts:
                        seen_texts.add(clean_txt)
                        items.append(current_item)

                clean_title = item_prefix_pattern.sub("", line).strip()
                clean_title = re.sub(r"^[:\.\-\s]+", "", clean_title)
                current_item = {
                    "raw_line": line,
                    "title": clean_title if len(clean_title) < 120 else clean_title[:117] + "...",
                    "text": line
                }
            elif current_item:
                # If short continuation or spec detail, append to current item
                if len(line) < 250 and not is_new_item_prefix:
                    current_item["text"] += " " + line
                else:
                    if len(current_item["text"]) > 10:
                        clean_txt = current_item["text"].strip()
                        if clean_txt not in seen_texts:
                            seen_texts.add(clean_txt)
                            items.append(current_item)
                    if has_equip or has_is:
                        clean_title = item_prefix_pattern.sub("", line).strip()
                        current_item = {
                            "raw_line": line,
                            "title": clean_title if len(clean_title) < 120 else clean_title[:117] + "...",
                            "text": line
                        }
                    else:
                        current_item = None

        if current_item and len(current_item["text"]) > 10:
            clean_txt = current_item["text"].strip()
            if clean_txt not in seen_texts:
                seen_texts.add(clean_txt)
                items.append(current_item)

        # Fallback if no structured numbered items were found:
        # scan for sentences or clauses that mention equipment keywords
        if not items:
            sentences = re.split(r"(?<=[.;\n])\s+", text)
            for s in sentences:
                s_clean = s.strip()
                if len(s_clean) > 15 and any(re.search(rf"\b{kw}\b", s_clean, re.IGNORECASE) for kw in self.equipment_keywords):
                    if s_clean not in seen_texts:
                        seen_texts.add(s_clean)
                        items.append({
                            "raw_line": s_clean,
                            "title": s_clean[:100] + ("..." if len(s_clean) > 100 else ""),
                            "text": s_clean
                        })
                if len(items) >= 20:
                    break

        return items

    def analyze_document(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """
        Full end-to-end document analysis:
        1. Extract text
        2. Segment technical sections
        3. Decompose into items
        4. Run recommendation, QCO compliance & alert evaluation for each item
        5. Aggregate audit findings
        """
        raw_text = self.extract_text(file_bytes, filename)
        if not raw_text:
            return {
                "success": False,
                "error": "The document contains no readable text or is empty.",
                "filename": filename
            }

        # Segment and decompose
        tech_chunks = self.segment_technical_sections(raw_text)
        filtered_text = "\n".join(tech_chunks)
        extracted_items = self.extract_discrete_items(filtered_text)

        if not extracted_items:
            # Fall back to trying the whole raw text
            extracted_items = self.extract_discrete_items(raw_text)

        evaluated_items = []
        mandatory_qco_count = 0
        superseded_count = 0
        compliant_count = 0

        # Pattern to detect cited standard in the officer's text
        cited_is_pattern = re.compile(r"\bIS\s*:?\s*(\d+)(?:\s*\(PART\s*\d+\))?", re.IGNORECASE)

        for idx, item in enumerate(extracted_items[:30]):  # Cap at 30 items per document for fast audit
            item_text = item["text"]
            
            # Find any explicit IS standard cited in the officer's document
            cited_matches = cited_is_pattern.findall(item_text)
            cited_standards = [f"IS {m}" for m in cited_matches] if cited_matches else []

            # Run engine recommendation
            rec = self.engine.process_query(item_text)
            if not rec.get("success"):
                continue

            primary_std = rec.get("primary_standard")
            qco = rec.get("qco", {})
            warnings = rec.get("warnings", [])
            certification = rec.get("certification", {})

            # Audit check: did officer cite an old or superseded standard?
            tender_standard_audit = {
                "cited_in_document": cited_standards,
                "audit_status": "OK",
                "audit_message": "Requirement is aligned with active BIS standards."
            }

            has_superseded_warning = False
            for w in warnings:
                level = str(w.get("level", "")).upper()
                badge = str(w.get("badge", "")).upper()
                title = str(w.get("title", "")).upper()
                if level == "RED" or "SUPERSEDED" in badge or "SUPERSEDED" in title or "OBSOLETE" in title:
                    has_superseded_warning = True
                    superseded_count += 1
                    tender_standard_audit["audit_status"] = "OBSOLETE_STANDARD"
                    tender_standard_audit["audit_message"] = w.get("message") or f"Standard cited is obsolete and superseded."
                    break

            if qco.get("status") == "Mandatory" or qco.get("applicable"):
                mandatory_qco_count += 1

            if not has_superseded_warning and primary_std:
                compliant_count += 1

            evaluated_items.append({
                "item_index": idx + 1,
                "title": item["title"],
                "specification_text": item_text,
                "detected_domains": rec.get("detected_domains", []),
                "extracted_parameters": rec.get("extracted_parameters", {}),
                "primary_standard": primary_std,
                "allied_standards": rec.get("allied_standards", {}),
                "certification": certification,
                "qco": qco,
                "warnings": warnings,
                "tender_standard_audit": tender_standard_audit,
                "recommended_tender_clause": rec.get("tender_clause", "")
            })

        total_items = len(evaluated_items)
        compliance_rate = round((compliant_count / total_items * 100), 1) if total_items > 0 else 100.0

        return {
            "success": True,
            "filename": filename,
            "file_size_bytes": len(file_bytes),
            "total_pages_approx": max(1, raw_text.count("--- PAGE ")),
            "summary": {
                "total_items_detected": total_items,
                "compliant_items": compliant_count,
                "superseded_flags": superseded_count,
                "mandatory_qco_items": mandatory_qco_count,
                "compliance_score_percent": compliance_rate,
                "executive_verdict": (
                    "CRITICAL ATTENTION REQUIRED: Obsolete/superseded standards detected in tender document."
                    if superseded_count > 0 else
                    "COMPLIANT: All identified equipment aligns with active BIS standards and QCO mandates."
                )
            },
            "evaluated_items": evaluated_items
        }
